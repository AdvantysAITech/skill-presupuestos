"""
Cuerpo del presupuesto de Vertical Projects con maqueta propia.

Vertical comparte el contenido con Scala (mismos apartados, mismos cálculos)
pero NO la maqueta, para que un cliente no pueda relacionar las dos empresas.
Paleta y tipografías: las del documento de referencia VRT-2026-0003 (grafito,
dorado y crema; Oswald para títulos y Lato para el texto).

Diferencias de composición respecto a Scala:
  - Rótulos de sección con el número grande en dorado y filete fino, no la
    barra con el número pequeño.
  - Tablas abiertas: sin bordes verticales, filetes horizontales finos y
    cabecera con línea dorada. Scala usa rejilla completa con cabecera gris.
  - Objeto y alcance en formato ficha (título a la izquierda, texto a la
    derecha).
  - Total destacado entre dos filetes dorados, sin barra de color.
  - Desglose detallado con cabecera de capítulo en grafito y subtotal en crema.
  - Garantías en tarjetas 2 x 2, condiciones generales numeradas, próximo paso
    con filete lateral y anexo fotográfico a dos columnas.
"""

from __future__ import annotations

import copy

from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm

from generar import (
    AVISOS, DECL, Constructor, D, cent, celda, el, fila, fmt, fmt_cant, fmt_cifra, fmt_pct, foto_comprimida, p,
    r, tabla,
)
from portada import a_frase

GRAFITO = "1C1917"
DORADO = "C9A45C"
DORADO_OSC = "9A7434"
CREMA = "F6F1E7"
TENUE = "7D756B"
FILETE = "D9D0C1"
OSWALD = "Oswald"
ANCHO_FOTO_ANEXO_CM, ALTO_FOTO_ANEXO_CM = 8.0, 6.0


def bordes(top=None, bottom=None, left=None, right=None) -> str:
    """Cada lado: None (sin borde) o (grosor en octavos de punto, color)."""
    def lado(nombre, v):
        if not v:
            return f'<w:{nombre} w:val="nil"/>'
        return f'<w:{nombre} w:val="single" w:sz="{v[0]}" w:space="0" w:color="{v[1]}"/>'
    return ("<w:tcBorders>" + lado("top", top) + lado("left", left) + lado("bottom", bottom)
            + lado("right", right) + "</w:tcBorders>")


SIN_BORDES = bordes()


def rotulo(texto, color=DORADO_OSC, tam=15, espaciado=40):
    return r(texto.upper(), negrita=True, tam=tam, color=color, espaciado=espaciado)


class ConstructorVertical(Constructor):

    _pegar_siguiente_tabla = False

    def add(self, *elementos):
        """Como el de Scala, pero la primera fila de la tabla que sigue a un rótulo
        se queda en la misma página que el rótulo."""
        for e in elementos:
            nuevos = el(e) if isinstance(e, str) else (e if isinstance(e, list) else [e])
            for x in nuevos:
                if self._pegar_siguiente_tabla and x.tag == qn("w:tbl"):
                    self._pegar_siguiente_tabla = False
                    filas = x.findall(qn("w:tr"))
                    if len(filas) > 1:
                        for par in filas[0].iter(qn("w:p")):
                            self._mantener_con_siguiente(par)
                self.salida.append(x)

    # --- rótulos ---------------------------------------------------------
    def seccion(self, titulo: str, salto=False):
        self.n_seccion += 1
        num = f"{self.n_seccion:02d}"
        a = [1100, self.ancho_tabla - 1100]
        bajo = (6, GRAFITO)
        izq = celda(p(r(num, negrita=True, tam=44, color=DORADO, fuente=OSWALD), despues=0, linea=240, keep_next=True),
                    a[0], bordes=bordes(bottom=bajo), alin_v="bottom", margen=(0, 0, 40, 0))
        der = celda(p(r(titulo.upper(), negrita=False, tam=24, color=GRAFITO, fuente=OSWALD, espaciado=30),
                      despues=0, linea=240, keep_next=True),
                    a[1], bordes=bordes(bottom=bajo), alin_v="bottom", margen=(0, 0, 70, 0))
        if salto:
            self.add(p("", salto=True, despues=0))
        self.add(p("", despues=0, linea=200, keep_next=True))
        self.add(tabla(a, [fila([izq, der])], margen_celda=0))
        self.add(p("", despues=0, linea=160, keep_next=True))
        self._pegar_siguiente_tabla = True
        return None

    def texto(self, texto, negrita=False, despues=100, keep_next=False, color=None, tam=18, cursiva=False):
        return p(r(texto, negrita=negrita, cursiva=cursiva, color=color or self.c["color_texto"], tam=tam),
                 despues=despues, linea=276, keep_next=keep_next)

    # --- apertura del cuerpo --------------------------------------------
    def cabecera(self):
        d, cli = self.datos, self.datos["cliente"]
        self.add(p("", despues=0))
        self.add(p(rotulo(f"Propuesta {d['numero']}  ·  Presupuesto y mediciones", tam=16, espaciado=50),
                   despues=40))
        self.add(p(r(d["titulo"], tam=56, color=GRAFITO, fuente=OSWALD), despues=200, linea=240))

        adm = (cli.get("administracion") or "").strip() or "—"
        if (cli.get("administracion_localidad") or "").strip():
            adm += f" · {cli['administracion_localidad'].strip()}"
        lugar = " · ".join(x for x in [cli.get("localidad"), cli.get("provincia")] if x)
        datos_fila = [
            ("Inmueble", cli["comunidad"], lugar),
            ("Administración", adm, ""),
            ("Emisión", self.fechas["emision"], f"Válido hasta {self.fechas['validez']}"),
        ]
        a = [3900, 3300, self.ancho_tabla - 7200]
        celdas = []
        for (etq, valor, detalle), w in zip(datos_fila, a):
            cont = p(rotulo(etq, tam=14), despues=40) + p(r(valor, negrita=True, tam=18, color=GRAFITO), despues=0)
            if detalle:
                cont += p(r(detalle, tam=15, color=TENUE), despues=0)
            celdas.append(celda(cont, w, bordes=bordes(top=(18, DORADO)), margen=(100, 0, 0, 160)))
        self.add(tabla(a, [fila(celdas)], margen_celda=0))

    # --- apartados --------------------------------------------------------
    def objeto(self):
        d = self.datos
        self.seccion("Objeto y alcance de la intervención")
        if d.get("objeto_intro"):
            self.add(self.texto(d["objeto_intro"]))
        items = [it if isinstance(it, dict) else {"titulo": "", "texto": it} for it in d.get("objeto") or []]
        if not items:
            return
        a = [2900, self.ancho_tabla - 2900]
        filas = []
        for i, it in enumerate(items, 1):
            linea = bordes(bottom=(4, FILETE))
            izq = (p(r(f"{i:02d}", negrita=True, tam=16, color=DORADO, fuente=OSWALD), despues=0)
                   + p(r((it.get("titulo") or "").upper(), tam=18, color=GRAFITO, fuente=OSWALD, espaciado=10),
                       despues=0, linea=240))
            der = p(r(it.get("texto", ""), tam=18, color=self.c["color_texto"]), despues=0, linea=276)
            filas.append(fila([celda(izq, a[0], bordes=linea, margen=(110, 0, 110, 120)),
                               celda(der, a[1], bordes=linea, margen=(110, 0, 110, 0))]))
        self.add(tabla(a, filas, margen_celda=0))

    def _cabecera_abierta(self, textos, anchos, numericas):
        return fila([celda(p(rotulo(t, color=GRAFITO, tam=14, espaciado=20), alin="right" if n else None,
                             keep_next=True), w,
                           bordes=bordes(bottom=(12, DORADO)), margen=(40, 80, 60, 80))
                     for t, w, n in zip(textos, anchos, numericas)], cabecera=True)

    def diagnostico(self):
        dg = self.datos.get("diagnostico") or {}
        pats = dg.get("patologias") or []
        if not pats:
            return
        self.seccion("Diagnóstico del edificio")
        if dg.get("intro"):
            self.add(self.texto(dg["intro"]))
        colores = {"urgente": "B00020", "grave": "B00020", "moderada": "B7791F", "media": "B7791F",
                   "baja": "2E7D32", "leve": "2E7D32"}
        a = [2700, self.ancho_tabla - 2700 - 1500, 1500]
        filas = [self._cabecera_abierta(["Zona / elemento", "Patología detectada", "Gravedad"], a, [False, False, True])]
        for pt in pats:
            grav = str(pt.get("gravedad", "")).strip()
            col = colores.get(grav.lower().strip("⚠● "), GRAFITO)
            linea = bordes(bottom=(4, FILETE))
            filas.append(fila([
                celda(p(r(pt.get("zona", ""), negrita=True, tam=17, color=GRAFITO), despues=0), a[0], bordes=linea,
                      margen=(80, 80, 80, 80)),
                celda(p(r(pt.get("descripcion", ""), tam=17), despues=0, linea=264), a[1], bordes=linea,
                      margen=(80, 80, 80, 80)),
                celda(p(r("■ ", tam=16, color=col) + r(grav.upper(), negrita=True, tam=15, color=col, espaciado=10),
                        alin="right", despues=0), a[2], bordes=linea, margen=(80, 80, 80, 80)),
            ]))
        self.add(tabla(a, filas, margen_celda=0))
        if dg.get("conclusion"):
            self.add(p("", despues=0), self.texto(dg["conclusion"]))

    def alertas(self):
        al = self.datos.get("alertas") or {}
        items = al.get("items") or []
        if not items or not al.get("incluir_en_doc"):
            return
        self.seccion("Alertas técnicas")
        ancho = self.ancho_tabla
        filas = []
        for it in items:
            cont = (p(r(it.get("titulo", "").upper(), tam=18, color=GRAFITO, fuente=OSWALD, espaciado=20), despues=40)
                    + p(r(it.get("descripcion", ""), tam=17), despues=0, linea=264))
            filas.append(fila([celda(cont, ancho, fondo=CREMA, bordes=bordes(left=(36, DORADO)),
                                     margen=(120, 220, 120, 220))]))
            filas.append(fila([celda(p("", despues=0, linea=80), ancho, bordes=SIN_BORDES)]))
        self.add(tabla([ancho], filas[:-1], margen_celda=0))

    def desglose_economico(self):
        calc = self.calc
        self.seccion("Resumen económico por capítulos")
        a = [1300, self.ancho_tabla - 1300 - 2400, 2400]
        filas = [self._cabecera_abierta(["Cap.", "Capítulo", "Importe"], a, [False, False, True])]
        for c in calc["capitulos"]:
            linea = bordes(bottom=(4, FILETE))
            filas.append(fila([
                celda(p(r(c["codigo"], negrita=True, tam=18, color=DORADO_OSC, fuente=OSWALD), despues=0), a[0],
                      bordes=linea, margen=(90, 80, 90, 80)),
                celda(p(r(a_frase(c["titulo"]), tam=18, color=GRAFITO), despues=0), a[1], bordes=linea,
                      margen=(90, 80, 90, 80)),
                celda(p(r(f"{fmt(c['subtotal'])} €", tam=18, color=GRAFITO), alin="right", despues=0), a[2],
                      bordes=linea, margen=(90, 80, 90, 80)),
            ]))
        self.add(tabla(a, filas, margen_celda=0))

        # Totales: columna derecha, sin barra de color.
        lineas = [("Presupuesto de ejecución material (PEM)", fmt(calc["pem"]))]
        if calc["margen_importe"]:
            lineas.append((f"Gastos generales y beneficio industrial ({fmt_pct(calc['margen_pct'])} %)",
                           fmt(calc["margen_importe"])))
            lineas.append(("Base imponible", fmt(calc["base"])))
        if calc["iva_pct"]:
            lineas.append((f"IVA {fmt_pct(calc['iva_pct'])} %", fmt(calc["iva"])))
        hueco = 3600
        a2 = [hueco, self.ancho_tabla - hueco - 2400, 2400]
        vacio = celda(p("", despues=0), a2[0], bordes=SIN_BORDES)
        filas = []
        for etq, valor in lineas:
            filas.append(fila([vacio,
                               celda(p(r(etq, tam=16, color=TENUE), despues=0), a2[1], bordes=SIN_BORDES,
                                     margen=(50, 80, 50, 80)),
                               celda(p(r(f"{valor} €", tam=17, color=GRAFITO), alin="right", despues=0), a2[2],
                                     bordes=SIN_BORDES, margen=(50, 80, 50, 80))]))
        etq_total = "Total presupuesto (sin IVA)" if calc["iva_pct"] == 0 else "Total presupuesto (IVA incluido)"
        marco = bordes(top=(18, DORADO), bottom=(18, DORADO))
        filas.append(fila([vacio,
                           celda(p(rotulo(etq_total, color=GRAFITO, tam=15, espaciado=30), despues=0), a2[1],
                                 bordes=marco, alin_v="center", margen=(140, 80, 140, 80)),
                           celda(p(r(f"{fmt(calc['total'])} €", negrita=True, tam=40, color=DORADO_OSC, fuente=OSWALD),
                                   alin="right", despues=0, linea=240), a2[2], bordes=marco, alin_v="center",
                                 margen=(100, 80, 100, 80))]))
        # Las líneas de totales no se separan del total al cambiar de página.
        self.add(p("", despues=0, linea=160, keep_next=True), self._junta(tabla(a2, filas, margen_celda=0)))

        if calc["pendientes"]:
            self.add(p(r("Las partidas marcadas como «Pendiente de precio» no están incluidas en el total. "
                         "Se valorarán y comunicarán a la comunidad antes de la aprobación del presupuesto.",
                         cursiva=True, tam=16, color=TENUE), antes=100, despues=60))

        vecinos = self.datos.get("num_vecinos")
        if vecinos:
            n = D(vecinos)
            por = cent(calc["total"] / n)
            linea = (f"Coste por propietario: {fmt(por)} € ({fmt_cifra(n)} propietarios, "
                     f"{'sin IVA' if calc['iva_pct'] == 0 else 'IVA incluido'}).")
            meses = self.datos.get("meses_financiacion")
            if meses:
                linea += f" Cuota orientativa: {fmt(cent(por / D(meses)))} € al mes durante {fmt_cifra(D(meses))} meses."
            self.add(p(r(linea, negrita=True, tam=17, color=GRAFITO), antes=120, despues=60))

        cg = self.datos.get("capitulo_gratis") or {}
        if cg.get("activo"):
            cap = next((c for c in calc["capitulos"] if c["codigo"] == str(cg.get("capitulo", "")).strip()), None)
            if cap is None:
                AVISOS.append("capitulo_gratis: no encuentro el capítulo indicado; no se incluye la oferta.")
            else:
                valor = f" (valor {fmt(cap['subtotal'])} €)" if cg.get("mostrar_valor", True) else ""
                texto = (f"Si la comunidad acepta todos los capítulos de este presupuesto, el capítulo {cap['codigo']} "
                         f"{a_frase(cap['titulo'])}{valor} se ejecuta sin coste. Este importe no está descontado del total.")
                ancho = self.ancho_tabla
                cont = p(rotulo("Oferta", tam=15), despues=40) + p(r(texto, tam=17, negrita=True, color=GRAFITO), despues=0)
                self.add(p("", despues=0), tabla([ancho], [fila([celda(cont, ancho, fondo=CREMA,
                                                                    bordes=bordes(left=(36, DORADO)),
                                                                    margen=(120, 220, 120, 220))])], 0))

    def condiciones_validez(self):
        self.seccion("Condiciones y validez")
        items = [(k, v.format(validez_dias=self.fechas["validez_dias"])) for k, v in self.marca["condiciones_resumen"]]
        if self.datos.get("plazo"):
            items.insert(2, ("PLAZO", f"{self.datos['plazo']} · orientativo"))
        a = [2200, self.ancho_tabla - 2200]
        filas = []
        for k, v in items:
            linea = bordes(bottom=(4, FILETE))
            filas.append(fila([celda(p(rotulo(k, tam=14), despues=0, keep_next=True), a[0], bordes=linea,
                                     margen=(70, 0, 70, 80)),
                               celda(p(r(v, tam=17, color=GRAFITO), despues=0, keep_next=True), a[1], bordes=linea,
                                     margen=(70, 80, 70, 0))]))
        self.add(tabla(a, filas, margen_celda=0), p("", despues=0, keep_next=True), self.clon(self.t_firmas1))

    def desglose_detallado(self):
        d = self.datos
        self.seccion("Mediciones y presupuesto detallado", salto=True)
        self.add(p(r(f"{d['titulo']}  ·  {d['cliente']['comunidad']}  ·  N/Ref. {d['numero']}", tam=16, color=TENUE),
                   despues=160))
        anchos = [1150, 4070, 560, 1000, 1300, 1640]
        escala = self.ancho_tabla / sum(anchos)
        anchos = [int(x * escala) for x in anchos]
        anchos[1] += self.ancho_tabla - sum(anchos)
        cabs = ["Código", "Descripción", "Ud", "Cant.", "Precio", "Importe"]
        num = [False, False, False, True, True, True]
        m = (60, 70, 60, 70)
        for c in self.calc["capitulos"]:
            filas = []
            titulo = (r(f"{c['codigo']}   ", negrita=True, tam=18, color=DORADO, fuente=OSWALD)
                      + r(c["titulo"].upper(), tam=18, color="FFFFFF", fuente=OSWALD, espaciado=20))
            filas.append(fila([
                celda(p(titulo, despues=0, keep_next=True), sum(anchos[:4]), fondo=GRAFITO, bordes=SIN_BORDES,
                      margen=(90, 140, 90, 70)),
                celda(p(r(f"{fmt(c['subtotal'])} €", negrita=True, tam=18, color=DORADO, fuente=OSWALD), alin="right",
                        despues=0, keep_next=True), sum(anchos[4:]), fondo=GRAFITO, bordes=SIN_BORDES,
                      margen=(90, 70, 90, 140)),
            ], cabecera=False))
            cab = self._cabecera_abierta(cabs, anchos, num)
            # La fila de capítulo ocupa dos celdas: rejilla propia con gridSpan.
            filas[0] = filas[0].replace(f'<w:tcW w:w="{sum(anchos[:4])}" w:type="dxa"/>',
                                        f'<w:tcW w:w="{sum(anchos[:4])}" w:type="dxa"/><w:gridSpan w:val="4"/>', 1)
            filas[0] = filas[0].replace(f'<w:tcW w:w="{sum(anchos[4:])}" w:type="dxa"/>',
                                        f'<w:tcW w:w="{sum(anchos[4:])}" w:type="dxa"/><w:gridSpan w:val="2"/>', 1)
            filas.append(cab)
            for pa in c["partidas"]:
                resumen = pa["resumen"]
                if pa["descripcion"] and pa["descripcion"] != resumen:
                    resumen = f"{resumen.rstrip('.')}. {pa['descripcion']}"
                vals = [pa["codigo"], resumen, pa["ud"], fmt_cant(pa["cantidad"]),
                        "—" if pa["pendiente"] else fmt(pa["precio"]),
                        "Pendiente de precio" if pa["pendiente"] else fmt(pa["importe"])]
                linea = bordes(bottom=(4, FILETE))
                celdas = []
                for i, (v, w, n) in enumerate(zip(vals, anchos, num)):
                    color = DORADO_OSC if i == 0 else GRAFITO
                    celdas.append(celda(p(r(v, tam=16, color=color, negrita=(i == 5 and pa["pendiente"])),
                                          alin="right" if n else None, despues=0, linea=252), w, bordes=linea,
                                        margen=m))
                filas.append(fila(celdas))
            filas.append(fila([
                celda(p(rotulo(f"Total capítulo {c['codigo']}", color=GRAFITO, tam=14, espaciado=20), despues=0),
                      sum(anchos[:5]), fondo=CREMA, bordes=SIN_BORDES, margen=(70, 70, 70, 140)),
                celda(p(r(f"{fmt(c['subtotal'])} €", negrita=True, tam=17, color=GRAFITO), alin="right", despues=0),
                      anchos[5], fondo=CREMA, bordes=SIN_BORDES, margen=(70, 70, 70, 70)),
            ]).replace(f'<w:tcW w:w="{sum(anchos[:5])}" w:type="dxa"/>',
                       f'<w:tcW w:w="{sum(anchos[:5])}" w:type="dxa"/><w:gridSpan w:val="5"/>', 1))
            self.add(tabla(anchos, filas, margen_celda=0), p("", despues=0, linea=280))

    def garantias(self):
        g = dict(self.marca["garantias"])
        g.update(self.datos.get("garantias") or {})
        # Sin salto de página: aprovecha el hueco que deja el desglose detallado.
        self.seccion(g.get("titulo") or "Garantías y certificaciones")
        if g.get("texto"):
            self.add(p(r(g["texto"], tam=30, color=DORADO_OSC, fuente=OSWALD), despues=200, linea=240))
        puntos = [list(x) for x in (g.get("puntos") or [])]
        for obra in self.datos.get("obras_referencia") or []:
            puntos.append([obra, ""])
        if not puntos:
            return
        hueco = 240
        ancho = (self.ancho_tabla - hueco) // 2
        a = [ancho, hueco, self.ancho_tabla - ancho - hueco]
        vacio = celda(p("", despues=0), hueco, bordes=SIN_BORDES)
        filas = []
        for i in range(0, len(puntos), 2):
            par = puntos[i:i + 2]
            celdas = []
            for j, (titulo, texto) in enumerate(par):
                cont = (p(r(f"{i + j + 1:02d}", negrita=True, tam=20, color=DORADO, fuente=OSWALD), despues=40)
                        + p(r(titulo.rstrip("."), tam=20, color=GRAFITO, fuente=OSWALD, espaciado=10), despues=60,
                            linea=240))
                if texto:
                    cont += p(r(texto, tam=16, color=self.c["color_texto"]), despues=0, linea=264)
                celdas.append(celda(cont, a[0] if j == 0 else a[2], fondo=CREMA, bordes=bordes(top=(24, DORADO)),
                                    margen=(160, 200, 180, 200)))
                if j == 0:
                    celdas.append(vacio)
            if len(par) == 1:
                celdas.append(celda(p("", despues=0), a[2], bordes=SIN_BORDES))
            filas.append(fila(celdas))
            filas.append(fila([celda(p("", despues=0, linea=120), w, bordes=SIN_BORDES) for w in a]))
        self.add(tabla(a, filas[:-1], margen_celda=0))

    def condiciones_generales(self):
        self.seccion("Condiciones generales")
        textos = [t.format(validez_dias=self.fechas["validez_dias"]) for t in self.marca["condiciones_generales"]]
        if self.datos.get("plazo"):
            textos.insert(2, f"Plazo de ejecución estimado: {self.datos['plazo']}. Es orientativo y puede ajustarse "
                             "según la meteorología, los accesos y lo que aparezca durante la obra.")
        textos += list(self.datos.get("condiciones_adicionales") or [])
        a = [800, self.ancho_tabla - 800]
        filas = []
        for i, t in enumerate(textos, 1):
            linea = bordes(bottom=(4, FILETE))
            filas.append(fila([
                celda(p(r(f"{i:02d}", negrita=True, tam=16, color=DORADO_OSC, fuente=OSWALD), despues=0), a[0],
                      bordes=linea, margen=(70, 0, 70, 0)),
                celda(p(r(t, tam=16, color=self.c["color_texto"]), despues=0, linea=264), a[1], bordes=linea,
                      margen=(70, 0, 70, 0)),
            ]))
        self.add(tabla(a, filas, margen_celda=0))

    def proximo_paso(self):
        pp = dict(self.marca["proximo_paso"])
        pp.update(self.datos.get("proximo_paso") or {})
        if not pp.get("titulo"):
            return
        ancho = self.ancho_tabla
        cont = (p(rotulo("Próximo paso", tam=14), despues=60)
                + p(r(pp["titulo"], tam=26, color=GRAFITO, fuente=OSWALD), despues=80, linea=240)
                + p(r(pp.get("texto", ""), tam=17, color=self.c["color_texto"]), despues=0, linea=276))
        self.add(p("", despues=160),
                 tabla([ancho], [fila([celda(cont, ancho, fondo=CREMA, bordes=bordes(left=(48, GRAFITO)),
                                             margen=(220, 300, 220, 300))])], 0),
                 p("", despues=160, keep_next=True))

    def _junta(self, xml_tabla: str):
        """Tabla cuyas filas no se separan entre páginas (todas salvo la última "con el siguiente")."""
        elementos = el(xml_tabla)
        filas = elementos[0].findall(qn("w:tr"))
        for tr in filas[:-1]:
            for par in tr.iter(qn("w:p")):
                self._mantener_con_siguiente(par)
        return elementos

    def firma_final(self):
        lugar = self.clon(self.p_lugar)
        self._mantener_con_siguiente(lugar)
        firmas = self.clon(self.t_firmas2)
        filas = firmas.findall(qn("w:tr"))
        for tr in filas[:-1]:
            for par in tr.iter(qn("w:p")):
                self._mantener_con_siguiente(par)
        self.add(lugar, firmas)

    def anexo(self, fotos):
        if not fotos:
            return
        lado, calidad = self.ajuste_foto
        self.seccion("Anexo fotográfico", salto=True)
        grupos: dict[str, list] = {}
        for f in fotos:
            grupos.setdefault(f.get("grupo") or "Estado actual", []).append(f)

        # Índice.
        a = [800, self.ancho_tabla - 800 - 2400, 2400]
        filas = []
        desde = 1
        for i, (nombre, lista) in enumerate(grupos.items(), 1):
            hasta = desde + len(lista) - 1
            rango = f"Foto {desde}" if len(lista) == 1 else f"Fotos {desde} a {hasta}"
            linea = bordes(bottom=(4, FILETE))
            filas.append(fila([
                celda(p(r(f"{i:02d}", negrita=True, tam=16, color=DORADO_OSC, fuente=OSWALD), despues=0), a[0],
                      bordes=linea, margen=(60, 0, 60, 0)),
                celda(p(r(nombre, tam=17, color=GRAFITO), despues=0), a[1], bordes=linea, margen=(60, 0, 60, 0)),
                celda(p(r(rango, tam=16, color=TENUE), alin="right", despues=0), a[2], bordes=linea,
                      margen=(60, 0, 60, 0)),
            ]))
            desde = hasta + 1
        self.add(tabla(a, filas, margen_celda=0))

        # Fotos a dos columnas.
        hueco = 280
        col = (self.ancho_tabla - hueco) // 2
        anchos = [col, hueco, self.ancho_tabla - col - hueco]
        numero = 0
        for i, (nombre, lista) in enumerate(grupos.items(), 1):
            self.add(p(r(f"{i:02d}   ", negrita=True, tam=22, color=DORADO, fuente=OSWALD)
                       + r(nombre.upper(), tam=22, color=GRAFITO, fuente=OSWALD, espaciado=20),
                       antes=360, despues=160, keep_next=True))
            for k in range(0, len(lista), 2):
                par = lista[k:k + 2]
                celdas = []
                for j, f in enumerate(par):
                    numero += 1
                    celdas.append(self._celda_foto(f, numero, anchos[0] if j == 0 else anchos[2], lado, calidad))
                    if j == 0:
                        celdas.append(celda(p("", despues=0), hueco, bordes=SIN_BORDES))
                if len(par) == 1:
                    celdas.append(celda(p("", despues=0), anchos[2], bordes=SIN_BORDES))
                self.add(tabla(anchos, [fila(celdas)], margen_celda=0), p("", despues=0, linea=200))

    def _celda_foto(self, f, numero, ancho, lado, calidad):
        buf, (w, h) = foto_comprimida(f["ruta"], lado, calidad)
        escala = min(ANCHO_FOTO_ANEXO_CM / w, ALTO_FOTO_ANEXO_CM / h)
        tmp = self.doc.add_paragraph()
        tmp.add_run().add_picture(buf, width=Cm(w * escala), height=Cm(h * escala))
        self.body.remove(tmp._p)
        tmp._p.insert(0, parse_xml(f'<w:pPr {DECL}><w:keepNext/><w:spacing w:before="0" w:after="80"/></w:pPr>'))
        pie = p(r(f"FOTO {numero:02d}", negrita=True, tam=14, color=DORADO_OSC, espaciado=30)
                + (r(f"   {f['pie']}", tam=15, color=TENUE) if f.get("pie") else ""), despues=0)
        c = parse_xml(f'<w:wrap {DECL}>{celda("", ancho, bordes=SIN_BORDES)}</w:wrap>').getchildren()[0]
        c.append(tmp._p)
        for e in parse_xml(f'<w:wrap {DECL}>{pie}</w:wrap>').getchildren():
            c.append(e)
        from lxml import etree
        return etree.tostring(c, encoding="unicode")
