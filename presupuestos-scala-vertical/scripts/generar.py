#!/usr/bin/env python3
"""
Genera el presupuesto en Word (.docx) con el diseño de Scala Valencia o de
Vertical Projects.

    python3 scripts/generar.py datos.json [salida.docx]

Estructura del documento:
  1. Portada infográfica a página completa (primera foto de la obra).
  2. Presupuesto y mediciones: objeto y alcance, diagnóstico, alertas técnicas,
     desglose económico por capítulos con totales, condiciones y validez, firmas.
  3. Desglose detallado: una tabla por capítulo.
  4. Garantías, condiciones generales, próximo paso y firmas.
  5. Anexo fotográfico con índice.

Los importes los calcula este script a partir de cantidad x precio (redondeo
a céntimos): no hace falta pasar importes ni totales en el JSON.
"""

from __future__ import annotations

import copy
import datetime as dt
import hashlib
import io
import json
import os
import re
import sys
import tempfile
from decimal import ROUND_HALF_UP, Decimal

try:
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn
    from docx.shared import Cm, Emu
    from PIL import Image, ImageOps
except ImportError:  # pragma: no cover
    sys.exit("Faltan dependencias: pip install python-docx pillow  (añade --break-system-packages si lo pide)")

try:  # fotos HEIC del iPhone
    import pillow_heif  # type: ignore

    pillow_heif.register_heif_opener()
except Exception:
    pass

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from portada import a_frase, generar_portada, generar_portada_vertical  # noqa: E402

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
DECL = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"'
)

LIMITE_BYTES = 3_600_000  # la app admite hasta ~4 MB por subida
AJUSTES_FOTO = [(1600, 80), (1400, 76), (1200, 72), (1000, 68), (850, 62), (700, 58)]
MAX_FOTOS = 40
ANCHO_FOTO_CM, ALTO_FOTO_CM = 16.0, 10.5
UNIDADES_MEDIBLES = {"m", "ml", "m²", "m2", "m³", "m3", "kg", "h", "día", "días", "mes"}

AVISOS: list[str] = []
# Ruta de cada foto del anexo -> SHA-1 de los bytes JPEG que se insertan en el .docx.
# Lo usa incrustar.py para saber en qué archivo de word/media/ ha acabado cada foto.
REGISTRO_FOTOS: dict[str, str] = {}


# ---------------------------------------------------------------------------
# Números
# ---------------------------------------------------------------------------

def D(v) -> Decimal:
    if v is None or v == "":
        return Decimal("0")
    if isinstance(v, str):
        v = v.replace("€", "").replace(" ", "").strip()
        if "," in v:
            v = v.replace(".", "").replace(",", ".")
    return Decimal(str(v))


def cent(v: Decimal) -> Decimal:
    return v.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def fmt(v: Decimal) -> str:
    """1234.5 -> '1.234,50'"""
    s = f"{cent(v):,.2f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def fmt_cant(v: Decimal) -> str:
    v = v.normalize()
    entero, _, dec = f"{v:f}".partition(".")
    signo = "-" if entero.startswith("-") else ""
    entero = entero.lstrip("-")
    miles = f"{int(entero):,}".replace(",", ".")
    dec = dec[:2]
    return f"{signo}{miles},{dec.ljust(2, '0')}"


def fmt_cifra(v: Decimal) -> str:
    """Cantidad sin ceros sobrantes: 84 -> '84', 12.5 -> '12,5'."""
    q = cent(v).normalize()
    entero, _, dec = f"{q:f}".partition(".")
    miles = f"{int(entero):,}".replace(",", ".")
    return f"{miles},{dec}" if dec else miles


def fmt_pct(v) -> str:
    v = D(v).normalize()
    return f"{v:f}".replace(".", ",")


def abreviar(v: Decimal) -> str:
    e = float(v)
    if e >= 1_000_000:
        return f"{round(e / 1_000_000, 2):.2f}".rstrip("0").rstrip(".").replace(".", ",") + "M €"
    if e >= 1000:
        return f"{round(e / 1000, 1):.1f}".replace(".", ",") + "k €"
    return f"{fmt(v)} €"


def repartir_porcentajes(importes: list[Decimal], total: Decimal) -> list[float]:
    if not importes:
        return []
    if total <= 0:
        return [0.0] * len(importes)
    bruto = [round(float(i) * 1000 / float(total)) / 10 for i in importes]
    mayor = max(range(len(importes)), key=lambda i: importes[i])
    resto = 1000 - sum(round(v * 10) for i, v in enumerate(bruto) if i != mayor)
    bruto[mayor] = resto / 10
    return bruto


# ---------------------------------------------------------------------------
# Cálculo
# ---------------------------------------------------------------------------

def calcular(datos: dict) -> dict:
    margen = datos.get("margen") or {}
    m_pct = D(margen.get("porcentaje", 0))
    base_margen = (margen.get("base") or "partida").lower()
    m_partida = m_pct if (m_pct and base_margen.startswith("partida")) else Decimal("0")
    m_total = m_pct if (m_pct and base_margen.startswith("total")) else Decimal("0")

    capitulos = []
    pendientes = 0
    for c in datos["capitulos"]:
        partidas = []
        subtotal = Decimal("0")
        for p in c.get("partidas", []):
            pendiente = bool(p.get("pendiente")) or p.get("precio") in (None, "", "pendiente", "Pendiente")
            cantidad = D(p.get("cantidad", p.get("medicion", 1)))
            if pendiente:
                precio = importe = None
                pendientes += 1
            else:
                precio = cent(D(p["precio"]) * (1 + m_partida / 100))
                importe = cent(cantidad * precio)
                subtotal += importe
            partidas.append(dict(
                codigo=str(p.get("codigo", "")).strip(),
                resumen=str(p.get("resumen") or p.get("titulo") or "").strip(),
                descripcion=str(p.get("descripcion") or "").strip(),
                ud=str(p.get("ud", "")).strip(),
                cantidad=cantidad, precio=precio, importe=importe, pendiente=pendiente,
            ))
        capitulos.append(dict(codigo=str(c.get("codigo") or c.get("num") or "").strip(),
                              titulo=str(c.get("titulo", "")).strip(), partidas=partidas, subtotal=subtotal))

    pem = sum((c["subtotal"] for c in capitulos), Decimal("0"))
    margen_importe = cent(pem * m_total / 100) if m_total else Decimal("0")
    base = pem + margen_importe
    iva_pct = D(datos.get("iva_porcentaje", datos.get("iva", 0)))
    iva = cent(base * iva_pct / 100)
    total = base + iva
    return dict(capitulos=capitulos, pem=pem, margen_pct=m_total, margen_importe=margen_importe, base=base,
                iva_pct=iva_pct, iva=iva, total=total, pendientes=pendientes)


# ---------------------------------------------------------------------------
# Portada
# ---------------------------------------------------------------------------

def datos_portada(datos, marca, calc, fechas, foto_portada):
    ident = marca["portada"]
    serie = ident["paleta"]["serie"]
    caps = calc["capitulos"]
    colores = [serie[i % len(serie)] for i in range(len(caps))]
    n = len(colores)
    if n > 2 and colores[-1] == colores[0]:
        libre = next((c for c in serie if c not in (colores[0], colores[-2])), None)
        if libre:
            colores[-1] = libre
    pcts = repartir_porcentajes([c["subtotal"] for c in caps], calc["pem"])
    capitulos = [dict(codigo=c["codigo"], nombre=c["titulo"], importe_fmt=f"{fmt(c['subtotal'])} €",
                      porcentaje=pcts[i], color=colores[i]) for i, c in enumerate(caps)]

    lineas = [(p, colores[i], i * 1000 + j) for i, c in enumerate(caps) for j, p in enumerate(c["partidas"])
              if not p["pendiente"]]
    lineas.sort(key=lambda t: (-t[0]["importe"], t[2]))
    mayor = lineas[0][0]["importe"] if lineas else Decimal("0")
    partidas = [dict(codigo=p["codigo"], resumen=a_frase(p["resumen"]), importe_fmt=f"{fmt(p['importe'])} €",
                     proporcion=float(p["importe"] / mayor) if mayor else 0, color=col) for p, col, _ in lineas[:5]]
    restantes = lineas[5:]
    partidas_restantes = (len(restantes), f"{fmt(sum((r[0]['importe'] for r in restantes), Decimal('0')))} €") if restantes else None

    cifras = []
    for p, _, _ in lineas:
        if len(cifras) >= 2:
            break
        ud = p["ud"].lower()
        if not (ud in UNIDADES_MEDIBLES or (ud in ("ud", "u", "uds") and p["cantidad"] > 1)):
            continue
        caja = dict(valor=f"{fmt_cifra(p['cantidad'])} {p['ud']}", etiqueta=a_frase(p["resumen"]))
        if caja not in cifras:
            cifras.append(caja)
    total_partidas = sum(len(c["partidas"]) for c in caps)
    relleno = [dict(valor=f"{total_partidas} {'partida' if total_partidas == 1 else 'partidas'}",
                    etiqueta=f"en {len(caps)} {'capítulo' if len(caps) == 1 else 'capítulos'}")]
    if datos.get("plazo"):
        relleno.append(dict(valor=str(datos["plazo"]), etiqueta="plazo estimado de ejecución"))
    relleno.append(dict(valor=f"{fechas['validez_dias']} días", etiqueta="validez del presupuesto"))
    if marca.get("garantia_cifra"):
        relleno.append(dict(valor=marca["garantia_cifra"], etiqueta="garantía de obra"))
    for caja in relleno:
        if len(cifras) >= 4:
            break
        cifras.append(caja)

    tipos = [t.strip() for t in datos.get("tipos_trabajo") or [] if str(t).strip()]
    tipos = list(dict.fromkeys(tipos))
    etiquetas = tipos[:3] + ([f"+{len(tipos) - 3}"] if len(tipos) > 3 else [])
    if not etiquetas:
        etiquetas = relleno[0]["valor"], f"{len(caps)} {'capítulo' if len(caps) == 1 else 'capítulos'}"

    cli = datos["cliente"]
    adm = (cli.get("administracion") or "").strip()
    if adm and (cli.get("administracion_localidad") or "").strip():
        adm = f"{adm} · {cli['administracion_localidad'].strip()}"

    sin_iva = calc["iva_pct"] == 0
    base_fmt = f"{fmt(calc['base'])} €"
    return dict(
        identidad=ident,
        antetitulo="PRESUPUESTO DE OBRA",
        titulo=datos["titulo"],
        comunidad=cli["comunidad"],
        localidad=cli.get("localidad", ""),
        numero=datos["numero"],
        fecha=fechas["emision"],
        fecha_validez=fechas["validez"],
        administrador=adm,
        foto_portada=foto_portada,
        etiquetas=list(etiquetas),
        capitulos=capitulos,
        partidas=partidas,
        partidas_restantes=partidas_restantes,
        cifras=cifras,
        pem_fmt=f"{fmt(calc['pem'])} €" if calc["margen_importe"] else base_fmt,
        pem_abreviado=abreviar(calc["pem"] if calc["margen_importe"] else calc["base"]),
        iva_etiqueta="Sin IVA" if sin_iva else f"IVA {fmt_pct(calc['iva_pct'])} %",
        iva_fmt=f"{fmt(calc['iva'])} €",
        total_etiqueta="TOTAL (SIN IVA)" if sin_iva else "TOTAL (IVA INCLUIDO)",
        total_fmt=f"{fmt(calc['total'])} €",
        base_fmt=base_fmt,
    )


# ---------------------------------------------------------------------------
# XML de Word
# ---------------------------------------------------------------------------

def esc(t) -> str:
    return (str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;"))


def rpr(negrita=False, cursiva=False, color=None, tam=None, fuente=None, espaciado=None, caps=False):
    x = ""
    if fuente:
        x += f'<w:rFonts w:ascii="{fuente}" w:hAnsi="{fuente}" w:eastAsia="{fuente}" w:cs="{fuente}"/>'
    if negrita:
        x += "<w:b/><w:bCs/>"
    if cursiva:
        x += "<w:i/><w:iCs/>"
    if caps:
        x += "<w:caps/>"
    if color:
        x += f'<w:color w:val="{color}"/>'
    if espaciado:
        x += f'<w:spacing w:val="{espaciado}"/>'
    if tam:
        x += f'<w:sz w:val="{tam}"/><w:szCs w:val="{tam}"/>'
    return f"<w:rPr>{x}</w:rPr>"


def r(texto, **kw) -> str:
    partes = str(texto).split("\n")
    cuerpo = '<w:br/>'.join(f'<w:t xml:space="preserve">{esc(p)}</w:t>' for p in partes)
    return f"<w:r>{rpr(**kw)}{cuerpo}</w:r>"


def ppr(alin=None, antes=0, despues=0, linea=None, keep_next=False, keep_lines=False, salto=False,
        borde_inf=None, sombreado=None, sangria=None):
    x = ""
    if keep_next:
        x += "<w:keepNext/>"
    if keep_lines:
        x += "<w:keepLines/>"
    if salto:
        x += "<w:pageBreakBefore/>"
    if borde_inf:
        x += f'<w:pBdr><w:bottom w:val="single" w:sz="6" w:space="3" w:color="{borde_inf}"/></w:pBdr>'
    if sombreado:
        x += f'<w:shd w:val="clear" w:color="auto" w:fill="{sombreado}"/>'
    lin = f' w:line="{linea}" w:lineRule="auto"' if linea else ""
    x += f'<w:spacing w:before="{antes}" w:after="{despues}"{lin}/>'
    if sangria:
        x += f'<w:ind w:left="{sangria[0]}" w:hanging="{sangria[1]}"/>'
    if alin:
        x += f'<w:jc w:val="{alin}"/>'
    return f"<w:pPr>{x}</w:pPr>"


def p(runs: str, **kw) -> str:
    return f"<w:p>{ppr(**kw)}{runs}</w:p>"


def el(xml: str):
    return parse_xml(f'<w:wrap {DECL}>{xml}</w:wrap>').getchildren()


def celda(contenido, ancho, fondo=None, borde="B8B8B8", alin_v="top", margen=None, bordes=None):
    tcb = bordes if bordes is not None else (
        f'<w:tcBorders><w:top w:val="single" w:sz="4" w:color="{borde}"/><w:left w:val="single" w:sz="4" w:color="{borde}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:color="{borde}"/><w:right w:val="single" w:sz="4" w:color="{borde}"/></w:tcBorders>'
        if borde else "")
    shd = f'<w:shd w:val="clear" w:color="auto" w:fill="{fondo}"/>' if fondo else ""
    mar = ""
    if margen:
        mar = (f'<w:tcMar><w:top w:w="{margen[0]}" w:type="dxa"/><w:left w:w="{margen[1]}" w:type="dxa"/>'
               f'<w:bottom w:w="{margen[2]}" w:type="dxa"/><w:right w:w="{margen[3]}" w:type="dxa"/></w:tcMar>')
    return (f'<w:tc><w:tcPr><w:tcW w:w="{ancho}" w:type="dxa"/>{tcb}{shd}{mar}<w:vAlign w:val="{alin_v}"/></w:tcPr>'
            f"{contenido}</w:tc>")


def fila(celdas: list[str], cabecera=False, no_partir=True) -> str:
    trpr = ("<w:cantSplit/>" if no_partir else "") + ("<w:tblHeader/>" if cabecera else "")
    return f"<w:tr><w:trPr>{trpr}</w:trPr>{''.join(celdas)}</w:tr>"


def tabla(anchos: list[int], filas: list[str], margen_celda=60) -> str:
    grid = "".join(f'<w:gridCol w:w="{a}"/>' for a in anchos)
    return (f'<w:tbl><w:tblPr><w:tblW w:w="{sum(anchos)}" w:type="dxa"/><w:tblLayout w:type="fixed"/>'
            f'<w:tblCellMar><w:top w:w="{margen_celda}" w:type="dxa"/><w:left w:w="80" w:type="dxa"/>'
            f'<w:bottom w:w="{margen_celda}" w:type="dxa"/><w:right w:w="80" w:type="dxa"/></w:tblCellMar>'
            f'<w:tblLook w:val="0000"/></w:tblPr><w:tblGrid>{grid}</w:tblGrid>{"".join(filas)}</w:tbl>')


def texto_de(elemento) -> str:
    return "".join(t.text or "" for t in elemento.iter(qn("w:t")))


def reemplazar_tokens(elemento, valores: dict):
    """Los tokens de la plantilla ocupan cada uno una run entera."""
    for t in elemento.iter(qn("w:t")):
        if t.text and "{{" in t.text:
            nuevo = t.text
            for k, v in valores.items():
                nuevo = nuevo.replace("{{" + k + "}}", v)
            t.text = nuevo
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def quitar_separadores_vacios(elemento):
    """'Nombre  ·  ' + '' -> 'Nombre'. Limpia separadores que se quedan colgando."""
    for par in elemento.iter(qn("w:p")):
        ts = [t for t in par.iter(qn("w:t"))]
        for i, t in enumerate(ts):
            txt = t.text or ""
            if txt.strip() in ("·", "(", ")", "") and txt != "":
                vecino_sig = ts[i + 1].text if i + 1 < len(ts) else ""
                if txt.strip() == "·" and not (vecino_sig or "").strip():
                    t.text = ""
                if txt.strip() == "(" and not (vecino_sig or "").strip():
                    t.text = ""
                    if i + 2 < len(ts) and (ts[i + 2].text or "").startswith(")"):
                        ts[i + 2].text = ts[i + 2].text[1:]


# ---------------------------------------------------------------------------
# Fotos
# ---------------------------------------------------------------------------

def foto_comprimida(ruta: str, lado: int, calidad: int):
    im = Image.open(ruta)
    im = ImageOps.exif_transpose(im).convert("RGB")
    im.thumbnail((lado, lado), Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=calidad, optimize=True, progressive=True)
    REGISTRO_FOTOS[ruta] = hashlib.sha1(buf.getvalue()).hexdigest()
    buf.seek(0)
    return buf, im.size


def foto_valida(ruta: str) -> bool:
    try:
        with Image.open(ruta) as im:
            im.verify()
        return True
    except Exception as e:  # noqa: BLE001
        AVISOS.append(f"No se ha podido abrir la foto {os.path.basename(ruta)} ({e}). Si es HEIC: pip install pillow-heif")
        return False


# ---------------------------------------------------------------------------
# Documento
# ---------------------------------------------------------------------------

class Constructor:
    def __init__(self, datos, marca, calc, fechas, ajuste_foto):
        self.datos, self.marca, self.calc, self.fechas = datos, marca, calc, fechas
        self.ajuste_foto = ajuste_foto
        self.es_vertical = datos["empresa"] == "vertical"
        self.estilo = marca["documento"]
        self.doc = Document(os.path.join(RAIZ, marca["plantilla"]))
        self.body = self.doc.element.body
        self.salida = []
        self.n_seccion = 0
        self._prototipos()

    # --- prototipos de la plantilla -------------------------------------
    def _prototipos(self):
        hijos = list(self.body.iterchildren())
        self.sect_final = hijos[-1]

        def buscar(pred, desde=0, tipo=None):
            for i, e in enumerate(hijos[desde:], desde):
                if tipo and e.tag != qn(tipo):
                    continue
                if pred(texto_de(e)):
                    return i, copy.deepcopy(e)
            raise RuntimeError("La plantilla no tiene el bloque esperado")

        self.p_portada = copy.deepcopy(hijos[0])
        _, self.p_ante = buscar(lambda t: t.strip() == "PRESUPUESTO Y MEDICIONES")
        i_tit, self.p_titulo = buscar(lambda t: t.strip() == "{{presup.TituloPresupuesto}}")
        _, self.t_inmueble = buscar(lambda t: "INMUEBLE" in t, tipo="w:tbl")
        if self.es_vertical:
            _, self.h_seccion = buscar(lambda t: "OBJETO Y ALCANCE" in t, tipo="w:p")
        else:
            _, self.h_seccion = buscar(lambda t: "OBJETO Y ALCANCE" in t, tipo="w:tbl")
        _, self.t_totales = buscar(lambda t: "PRESUPUESTO DE EJECUCIÓN MATERIAL (PEM)" in t, tipo="w:tbl")
        _, self.t_condiciones = buscar(lambda t: "VALIDEZ:" in t, tipo="w:tbl")
        _, self.t_firmas1 = buscar(lambda t: "Por la parte contratante" in t and "Fdo" not in t, tipo="w:tbl")
        i_da, self.p_desg_ante = buscar(lambda t: "DESGLOSE DETALLADO" in t)
        _, self.p_desg_titulo = buscar(lambda t: t.strip() == "{{presup.TituloPresupuesto}}", desde=i_da)
        _, self.p_desg_ref = buscar(lambda t: t.startswith("N/Ref:"))
        _, self.p_desg_cliente = buscar(lambda t: t.startswith("{{presup.ComunidadNombre}}"), desde=i_da)
        _, self.p_vineta = buscar(lambda t: t.startswith("Validez de la oferta"))
        _, self.p_lugar = buscar(lambda t: t.startswith("En {{doc.Localidad}}"))
        _, self.t_firmas2 = buscar(lambda t: "Fdo" in t, tipo="w:tbl")
        self.ancho_tabla = int(self.t_totales.find(qn("w:tblPr")).find(qn("w:tblW")).get(qn("w:w")))

    # --- utilidades ------------------------------------------------------
    @property
    def c(self):
        return self.estilo

    def add(self, *elementos):
        for e in elementos:
            if isinstance(e, str):
                self.salida.extend(el(e))
            elif isinstance(e, list):
                self.salida.extend(e)
            else:
                self.salida.append(e)

    def tokens(self):
        d, cli = self.datos, self.datos["cliente"]
        return {
            "presup.TituloPresupuesto": d["titulo"],
            "presup.ComunidadNombre": cli["comunidad"],
            "presup.ComunidadLocalidad": cli.get("localidad", ""),
            "presup.ComunidadProvincia": cli.get("provincia", "") or "",
            "presup.AdministradorNombre": (cli.get("administracion") or "—").strip(),
            "presup.AdministradorLocalidad": (cli.get("administracion_localidad") or "").strip(),
            "doc.NRef": d["numero"],
            "doc.FechaEmision": self.fechas["emision"],
            "doc.FechaValidez": self.fechas["validez"],
            "doc.Localidad": d.get("lugar_firma") or "Valencia",
        }

    def clon(self, proto, tokens=True):
        e = copy.deepcopy(proto)
        if tokens:
            reemplazar_tokens(e, self.tokens())
            quitar_separadores_vacios(e)
        return e

    def parrafo_texto(self, texto, negrita=False, despues=100, keep_next=False, color=None):
        tam = 18
        return p(r(texto, negrita=negrita, color=color or self.c["color_texto"], tam=tam),
                 despues=despues, linea=276 if self.es_vertical else 259, keep_next=keep_next)

    def seccion(self, titulo: str, salto=False):
        self.n_seccion += 1
        num = f"{self.n_seccion:02d}"
        h = copy.deepcopy(self.h_seccion)
        ts = list(h.iter(qn("w:t")))
        if self.es_vertical:
            ts[0].text = num
            ts[1].text = f"   {titulo.upper()}"
            ts[1].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            destino = h
        else:
            ts[-1].text = f"{num}  {titulo.upper()}"
            destino = None
        if salto:
            if self.es_vertical:
                self._salto_antes(h)
            else:
                self.add(p("", salto=True, despues=0))
        if not self.es_vertical:
            # El rótulo es una tabla: se marca "mantener con el siguiente" en sus
            # párrafos para que no se quede solo al pie de la página.
            for par in h.iter(qn("w:p")):
                self._mantener_con_siguiente(par)
            self.add(p("", despues=0, linea=120))
        self.add(h)
        self.add(p("", despues=0, linea=120, keep_next=True) if not self.es_vertical else [])
        return destino

    @staticmethod
    def _mantener_con_siguiente(par):
        ppr_el = par.find(qn("w:pPr"))
        if ppr_el is None:
            ppr_el = parse_xml(f'<w:pPr {DECL}/>')
            par.insert(0, ppr_el)
        if ppr_el.find(qn("w:keepNext")) is not None:
            return
        nuevo = parse_xml(f'<w:keepNext {DECL}/>')
        estilo = ppr_el.find(qn("w:pStyle"))
        if estilo is not None:
            estilo.addnext(nuevo)
        else:
            ppr_el.insert(0, nuevo)

    @staticmethod
    def _salto_antes(par):
        ppr_el = par.find(qn("w:pPr"))
        nuevo = parse_xml(f'<w:pageBreakBefore {DECL}/>')
        keep = ppr_el.find(qn("w:keepNext"))
        (keep if keep is not None else ppr_el.find(qn("w:pStyle"))).addnext(nuevo)

    # --- bloques ---------------------------------------------------------
    def portada(self, ruta_portada):
        par = copy.deepcopy(self.p_portada)
        for run in par.findall(qn("w:r")):
            par.remove(run)
        # LibreOffice deja el sectPr al principio del pPr; Word lo exige al final.
        ppr_el = par.find(qn("w:pPr"))
        sect = ppr_el.find(qn("w:sectPr"))
        if sect is not None:
            ppr_el.remove(sect)
            rpr_el = ppr_el.find(qn("w:rPr"))
            if rpr_el is not None:
                rpr_el.addnext(sect)
            else:
                ppr_el.append(sect)
        tmp = self.doc.add_paragraph()
        corrida = tmp.add_run()
        corrida.add_picture(ruta_portada, width=Emu(7560310), height=Emu(10692130))
        inline = corrida._r.find(".//" + qn("wp:inline"))
        grafico = inline.find(qn("a:graphic"))
        anclaje = parse_xml(
            f'<wp:anchor {DECL} distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" '
            f'behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
            f'<wp:simplePos x="0" y="0"/>'
            f'<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
            f'<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
            f'<wp:extent cx="7560310" cy="10692130"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
            f'<wp:wrapNone/><wp:docPr id="9001" name="Portada"/><wp:cNvGraphicFramePr/></wp:anchor>'
        )
        anclaje.append(grafico)
        drawing = corrida._r.find(qn("w:drawing"))
        drawing.remove(inline)
        drawing.append(anclaje)
        par.append(corrida._r)
        self.body.remove(tmp._p)
        self.add(par)

    def cabecera(self):
        self.add(p("", despues=0))
        self.add(self.clon(self.p_ante), self.clon(self.p_titulo), self.clon(self.t_inmueble))

    def objeto(self):
        d = self.datos
        self.seccion("Objeto y alcance de la intervención")
        if d.get("objeto_intro"):
            self.add(self.parrafo_texto(d["objeto_intro"]))
        for i, it in enumerate(d.get("objeto") or [], 1):
            if isinstance(it, str):
                self.add(self.parrafo_texto(it))
                continue
            if it.get("titulo"):
                self.add(self.parrafo_texto(f"{i:02d} · {it['titulo'].upper()}", negrita=True, despues=20, keep_next=True))
            if it.get("texto"):
                self.add(self.parrafo_texto(it["texto"]))

    def diagnostico(self):
        dg = self.datos.get("diagnostico") or {}
        pats = dg.get("patologias") or []
        if not pats:
            return
        self.seccion("Diagnóstico del edificio")
        if dg.get("intro"):
            self.add(self.parrafo_texto(dg["intro"]))
        colores = {"urgente": "B00020", "grave": "B00020", "moderada": "B7791F", "media": "B7791F", "baja": "2E7D32", "leve": "2E7D32"}
        anchos = [500, 2500, 5000, 1360]
        anchos[2] = self.ancho_tabla - sum(anchos) + anchos[2]
        cab = fila([celda(p(r(t, negrita=True, tam=16), alin=a), w, fondo=self.c["fondo_cabecera_tabla"])
                    for t, w, a in zip(["Nº", "ZONA / ELEMENTO", "PATOLOGÍA DETECTADA", "GRAVEDAD"], anchos,
                                       ["left", "left", "left", "center"])], cabecera=True)
        filas = [cab]
        for i, pt in enumerate(pats, 1):
            grav = str(pt.get("gravedad", "")).strip()
            col = colores.get(grav.lower().strip("⚠● "), self.c["color_texto"])
            filas.append(fila([
                celda(p(r(f"{i:02d}", negrita=True, tam=16)), anchos[0]),
                celda(p(r(pt.get("zona", ""), negrita=True, tam=16)), anchos[1]),
                celda(p(r(pt.get("descripcion", ""), tam=16)), anchos[2]),
                celda(p(r(grav.upper(), negrita=True, tam=15, color=col), alin="center"), anchos[3]),
            ]))
        self.add(tabla(anchos, filas))
        if dg.get("conclusion"):
            self.add(p("", despues=0), self.parrafo_texto(dg["conclusion"]))

    def alertas(self):
        al = self.datos.get("alertas") or {}
        items = al.get("items") or []
        if not items or not al.get("incluir_en_doc"):
            return
        self.seccion("Alertas técnicas")
        ancho = self.ancho_tabla
        bordes = ('<w:tcBorders><w:top w:val="nil"/><w:left w:val="single" w:sz="24" w:color="D98E04"/>'
                  '<w:bottom w:val="nil"/><w:right w:val="nil"/></w:tcBorders>')
        filas = []
        for it in items:
            contenido = (p(r(f"⚠  {it.get('titulo', '').upper()}", negrita=True, tam=16, color="8A5A00"), despues=30)
                         + p(r(it.get("descripcion", ""), tam=16), despues=0, linea=259))
            filas.append(fila([celda(contenido, ancho, fondo="FFF7E6", bordes=bordes, margen=(90, 160, 90, 160))]))
            filas.append(fila([celda(p("", despues=0, linea=60), ancho, bordes='<w:tcBorders><w:top w:val="nil"/><w:left w:val="nil"/><w:bottom w:val="nil"/><w:right w:val="nil"/></w:tcBorders>')]))
        self.add(tabla([ancho], filas[:-1], margen_celda=0))

    def desglose_economico(self):
        calc = self.calc
        self.seccion("Desglose económico por capítulos")
        a1 = int(self.ancho_tabla * 0.72)
        anchos = [a1, self.ancho_tabla - a1]
        filas = [fila([celda(p(r("CAPÍTULO", negrita=True, tam=16)), anchos[0], fondo=self.c["fondo_cabecera_tabla"]),
                       celda(p(r("IMPORTE", negrita=True, tam=16), alin="right"), anchos[1], fondo=self.c["fondo_cabecera_tabla"])],
                      cabecera=True)]
        for c in calc["capitulos"]:
            filas.append(fila([celda(p(r(f"{c['codigo']}  {c['titulo'].upper()}", tam=16)), anchos[0]),
                               celda(p(r(f"{fmt(c['subtotal'])} €", tam=16), alin="right"), anchos[1])]))
        self.add(tabla(anchos, filas))
        self.add(p("", despues=0))

        t = self.clon(self.t_totales, tokens=False)
        filas_t = t.findall(qn("w:tr"))
        valores = {
            "presup.TotalPEM": fmt(calc["pem"]),
            "presup.IVAPorcentaje": fmt_pct(calc["iva_pct"]),
            "presup.IVAImporte": fmt(calc["iva"]),
            "presup.TotalConIVA": fmt(calc["total"]),
        }
        reemplazar_tokens(t, valores)
        if calc["margen_importe"]:
            nueva = copy.deepcopy(filas_t[1])
            celdas_n = nueva.findall(qn("w:tc"))
            for i_c, (celda_n, valor) in enumerate(zip(
                    celdas_n, [f"GASTOS GENERALES Y BENEFICIO INDUSTRIAL ({fmt_pct(calc['margen_pct'])} %)",
                               fmt(calc["margen_importe"])])):
                ts = list(celda_n.iter(qn("w:t")))
                if i_c == 0:
                    ts[0].text = valor
                    for x in ts[1:]:
                        x.text = ""
                elif len(ts) > 1:
                    ts[0].text = valor
                else:
                    ts[0].text = f"{valor}  €"
            filas_t[1].addprevious(nueva)
        if calc["iva_pct"] == 0:
            t.remove(filas_t[1])
            for x in filas_t[2].iter(qn("w:t")):
                if "IVA INCLUIDO" in (x.text or ""):
                    x.text = x.text.replace("IVA INCLUIDO", "SIN IVA")
        self.add(t)

        if calc["pendientes"]:
            self.add(p(r("Las partidas marcadas como «Pendiente de precio» no están incluidas en el total. "
                         "Se valorarán y comunicarán a la comunidad antes de la aprobación del presupuesto.",
                         cursiva=True, tam=16, color=self.c["color_tenue"]), antes=60, despues=60))

        vecinos = self.datos.get("num_vecinos")
        if vecinos:
            n = D(vecinos)
            por = cent(calc["total"] / n)
            linea = f"Coste por propietario: {fmt(por)} € ({fmt_cifra(n)} propietarios, {'sin IVA' if calc['iva_pct'] == 0 else 'IVA incluido'})."
            meses = self.datos.get("meses_financiacion")
            if meses:
                linea += f" Cuota orientativa: {fmt(cent(por / D(meses)))} € al mes durante {fmt_cifra(D(meses))} meses."
            self.add(p(r(linea, negrita=True, tam=17, color=self.c["color_principal"]), antes=100, despues=60))

        cg = self.datos.get("capitulo_gratis") or {}
        if cg.get("activo"):
            cap = next((c for c in calc["capitulos"] if c["codigo"] == str(cg.get("capitulo", "")).strip()), None)
            if cap is None:
                AVISOS.append("capitulo_gratis: no encuentro el capítulo indicado; no se incluye la oferta.")
            else:
                valor = f" (valor {fmt(cap['subtotal'])} €)" if cg.get("mostrar_valor", True) else ""
                texto = (f"Si la comunidad acepta todos los capítulos de este presupuesto, el capítulo {cap['codigo']} "
                         f"{a_frase(cap['titulo'])}{valor} se ejecuta sin coste. Este importe no está descontado del total.")
                ancho = self.ancho_tabla
                contenido = (p(r("OFERTA", negrita=True, tam=15, espaciado=40, color=self.c["color_acento"]), despues=30)
                             + p(r(texto, tam=17, negrita=True, color=self.c["color_principal"]), despues=0))
                self.add(p("", despues=0), tabla([ancho], [fila([celda(contenido, ancho, fondo=self.c["fondo_suave"],
                                                                   borde=self.c["color_acento"], margen=(100, 160, 100, 160))])], 0))

    def condiciones_validez(self):
        self.seccion("Condiciones y validez")
        t = self.clon(self.t_condiciones, tokens=False)
        items = [(k, v.format(validez_dias=self.fechas["validez_dias"])) for k, v in self.marca["condiciones_resumen"]]
        if self.datos.get("plazo"):
            items.insert(2, ("PLAZO", f"{self.datos['plazo']} · orientativo"))
        mitad = -(-len(items) // 2)
        celdas = t.findall(".//" + qn("w:tc"))
        for celda_el, grupo in zip(celdas, [items[:mitad], items[mitad:]]):
            pars = celda_el.findall(qn("w:p"))
            proto = next((x for x in pars if ":" in texto_de(x)), pars[0])
            for x in pars:
                celda_el.remove(x)
            for k, v in grupo:
                nuevo = copy.deepcopy(proto)
                ts = list(nuevo.iter(qn("w:t")))
                ts[0].text = f"{k}:  "
                ts[-1].text = v
                for x in ts[1:-1]:
                    x.text = ""
                for x in ts:
                    x.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                celda_el.append(nuevo)
        for par in t.iter(qn("w:p")):
            self._mantener_con_siguiente(par)
        self.add(t, p("", despues=0, keep_next=True), self.clon(self.t_firmas1))

    def desglose_detallado(self):
        ante = self.clon(self.p_desg_ante)
        self._salto_antes(ante)
        self.add(ante, self.clon(self.p_desg_titulo), self.clon(self.p_desg_ref), self.clon(self.p_desg_cliente))
        anchos = [1224, 3672, 576, 1008, 1296, 1584]
        escala = self.ancho_tabla / sum(anchos)
        anchos = [int(a * escala) for a in anchos]
        anchos[1] += self.ancho_tabla - sum(anchos)
        cabs = ["CÓDIGO", "RESUMEN", "UD", "CANT.", "PRECIO", "IMPORTE"]
        num = [False, False, False, True, True, True]
        fondo_cab, fondo_tot = self.c["fondo_cabecera_tabla"], self.c["fondo_total_tabla"]
        for c in self.calc["capitulos"]:
            self.add(p(r(f"{c['codigo']}  {c['titulo'].upper()}", negrita=True, tam=20,
                         color=self.c["color_principal"]), antes=200, despues=60, keep_next=True))
            filas = [fila([celda(p(r(t, negrita=True, tam=16), alin="right" if n else None), w, fondo=fondo_cab)
                           for t, w, n in zip(cabs, anchos, num)], cabecera=True)]
            for pa in c["partidas"]:
                resumen = pa["resumen"]
                if pa["descripcion"] and pa["descripcion"] != resumen:
                    resumen = f"{resumen.rstrip('.')}. {pa['descripcion']}"
                vals = [pa["codigo"], resumen, pa["ud"], fmt_cant(pa["cantidad"]),
                        "—" if pa["pendiente"] else fmt(pa["precio"]),
                        "Pendiente de precio" if pa["pendiente"] else fmt(pa["importe"])]
                filas.append(fila([celda(p(r(v, tam=16, negrita=(i == 5 and pa["pendiente"])),
                                           alin="right" if n else None, linea=240), w)
                                   for i, (v, w, n) in enumerate(zip(vals, anchos, num))]))
            filas.append(fila([celda(p(r(v, negrita=True, tam=16), alin="right" if n else None), w, fondo=fondo_tot)
                               for v, w, n in zip(["", f"TOTAL {c['codigo']}", "", "", "", fmt(c["subtotal"])], anchos, num)]))
            self.add(tabla(anchos, filas, margen_celda=40))

    def garantias(self):
        g = dict(self.marca["garantias"])
        g.update(self.datos.get("garantias") or {})
        self.seccion("Garantías y certificaciones", salto=True)
        if g.get("titulo"):
            self.add(self.parrafo_texto(g["titulo"], negrita=True, despues=40, color=self.c["color_principal"]))
        if g.get("texto"):
            self.add(self.parrafo_texto(g["texto"]))
        puntos = list(g.get("puntos") or [])
        for obra in self.datos.get("obras_referencia") or []:
            puntos.append([obra, ""])
        for titulo, texto in puntos:
            runs = r("✓  ", negrita=True, tam=18, color=self.c["color_acento"]) + r(titulo, negrita=True, tam=18)
            if texto:
                runs += r(f"  {texto}", tam=18)
            self.add(p(runs, despues=90, linea=276, sangria=(300, 300)))

    def condiciones_generales(self):
        self.seccion("Condiciones generales")
        textos = [t.format(validez_dias=self.fechas["validez_dias"]) for t in self.marca["condiciones_generales"]]
        if self.datos.get("plazo"):
            textos.insert(2, f"Plazo de ejecución estimado: {self.datos['plazo']}. Es orientativo y puede ajustarse según "
                             "la meteorología, los accesos y lo que aparezca durante la obra.")
        textos += list(self.datos.get("condiciones_adicionales") or [])
        for t in textos:
            v = copy.deepcopy(self.p_vineta)
            ts = list(v.iter(qn("w:t")))
            ts[0].text = t
            for x in ts[1:]:
                x.text = ""
            self.add(v)

    def proximo_paso(self):
        pp = dict(self.marca["proximo_paso"])
        pp.update(self.datos.get("proximo_paso") or {})
        if not pp.get("titulo"):
            return
        ancho = self.ancho_tabla
        oscuro = self.c["color_principal"]
        contenido = (p(r("PRÓXIMO PASO", negrita=True, tam=15, espaciado=60, color=self.c.get("color_rotulo_sobre_oscuro", self.c["color_acento"])), alin="center", despues=60)
                     + p(r(pp["titulo"], negrita=True, tam=24, color="FFFFFF",
                           fuente=self.c.get("fuente_titulos")), alin="center", despues=80)
                     + p(r(pp.get("texto", ""), tam=17, color="E8E8E8"), alin="center", despues=0, linea=276))
        self.add(p("", despues=120),
                 tabla([ancho], [fila([celda(contenido, ancho, fondo=oscuro, borde=None, margen=(220, 300, 220, 300))])], 0),
                 p("", despues=120))

    def firma_final(self):
        self.add(self.clon(self.p_lugar), self.clon(self.t_firmas2))

    def anexo(self, fotos):
        if not fotos:
            return
        lado, calidad = self.ajuste_foto
        fuente_tit = self.c.get("fuente_titulos")
        self.add(p(r("Anexo fotográfico", negrita=True, tam=32, fuente=fuente_tit, color=self.c["color_principal"]),
                   salto=True, despues=200, keep_next=True))
        grupos: dict[str, list] = {}
        for f in fotos:
            grupos.setdefault(f.get("grupo") or "Estado actual", []).append(f)
        desde = 1
        for nombre, lista in grupos.items():
            hasta = desde + len(lista) - 1
            rango = f"foto {desde}" if len(lista) == 1 else f"fotos {desde} a {hasta}"
            self.add(p(r(f"{nombre} — {rango}", tam=20), despues=60, sangria=(280, 0)))
            desde = hasta + 1
        numero = 0
        for nombre, lista in grupos.items():
            self.add(p(r(nombre, negrita=True, tam=24, color=self.c["color_principal"]), antes=300, despues=150,
                       keep_next=True, borde_inf=self.c["borde_tabla"]))
            for f in lista:
                numero += 1
                buf, (w, h) = foto_comprimida(f["ruta"], lado, calidad)
                escala = min(ANCHO_FOTO_CM / w, ALTO_FOTO_CM / h)
                tmp = self.doc.add_paragraph()
                tmp.add_run().add_picture(buf, width=Cm(w * escala), height=Cm(h * escala))
                tmp_ppr = parse_xml(f'<w:pPr {DECL}><w:keepNext/><w:spacing w:before="160" w:after="60"/><w:jc w:val="center"/></w:pPr>')
                tmp._p.insert(0, tmp_ppr)
                self.body.remove(tmp._p)
                self.add(tmp._p)
                pie = f"Foto {numero}" + (f". {f['pie']}" if f.get("pie") else "")
                self.add(p(r(pie, cursiva=True, tam=18, color=self.c["color_tenue"]), alin="center", despues=200))

    # --- montaje -------------------------------------------------------------
    def montar(self, ruta_portada, fotos):
        self.portada(ruta_portada)
        self.cabecera()
        self.objeto()
        self.diagnostico()
        self.alertas()
        self.desglose_economico()
        self.condiciones_validez()
        self.desglose_detallado()
        self.garantias()
        self.condiciones_generales()
        self.proximo_paso()
        self.firma_final()
        self.anexo(fotos)

        for hijo in list(self.body.iterchildren()):
            if hijo is not self.sect_final:
                self.body.remove(hijo)
        for e in self.salida:
            self.sect_final.addprevious(e)

        # Cabeceras y pies de la sección de contenido.
        for s in self.doc.sections:
            for parte in (s.header, s.footer, s.first_page_header, s.first_page_footer):
                try:
                    if not parte.is_linked_to_previous:
                        reemplazar_tokens(parte._element, self.tokens())
                except Exception:
                    pass

        props = self.doc.core_properties
        props.title = f"Presupuesto {self.datos['numero']} - {self.datos['titulo']}"
        props.author = self.marca["empresa"]["razon_social"]
        props.comments = ""
        return self.doc


# ---------------------------------------------------------------------------
# Entrada
# ---------------------------------------------------------------------------

def validar(datos: dict):
    faltan = [k for k in ("empresa", "numero", "titulo", "cliente", "capitulos") if not datos.get(k)]
    if faltan:
        sys.exit(f"ERROR: faltan campos obligatorios en el JSON: {', '.join(faltan)}")
    if datos["empresa"] not in ("scala", "vertical"):
        sys.exit('ERROR: "empresa" tiene que ser "scala" o "vertical"')
    if not datos["cliente"].get("comunidad"):
        sys.exit("ERROR: falta cliente.comunidad")
    for c in datos["capitulos"]:
        if not c.get("partidas"):
            sys.exit(f"ERROR: el capítulo {c.get('codigo')} no tiene partidas")
        for pa in c["partidas"]:
            if "cantidad" not in pa and "medicion" not in pa:
                sys.exit(f"ERROR: la partida {pa.get('codigo')} no tiene cantidad")
    if datos.get("cliente", {}).get("cif"):
        AVISOS.append("Se ha ignorado cliente.cif: el CIF de la comunidad no va en el documento.")


def fechas_de(datos, marca):
    emision = datos.get("fecha") or dt.date.today().strftime("%d/%m/%Y")
    dias = int(datos.get("validez_dias") or marca["validez_dias"])
    try:
        f = dt.datetime.strptime(emision, "%d/%m/%Y").date()
    except ValueError:
        sys.exit('ERROR: "fecha" tiene que ir como dd/mm/aaaa')
    return dict(emision=f.strftime("%d/%m/%Y"), validez=(f + dt.timedelta(days=dias)).strftime("%d/%m/%Y"),
                validez_dias=dias)


def ruta_salida_por_defecto(datos):
    base = "/mnt/user-data/outputs" if os.path.isdir("/mnt/user-data/outputs") else os.getcwd()
    nombre = datos.get("nombre_archivo") or f"Presupuesto {datos['numero']} - {datos['cliente']['comunidad']}"
    nombre = re.sub(r'[\\/:*?"<>|]+', " ", nombre).strip()
    if not nombre.lower().endswith(".docx"):
        nombre += ".docx"
    return os.path.join(base, nombre)


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    with open(sys.argv[1], encoding="utf-8") as fh:
        datos = json.load(fh)
    validar(datos)
    with open(os.path.join(RAIZ, "assets", "marcas.json"), encoding="utf-8") as fh:
        marca = json.load(fh)[datos["empresa"]]
    salida = sys.argv[2] if len(sys.argv) > 2 else ruta_salida_por_defecto(datos)

    calc = calcular(datos)
    fechas = fechas_de(datos, marca)

    fotos = []
    for f in datos.get("fotos") or []:
        f = {"ruta": f} if isinstance(f, str) else dict(f)
        if f.get("ruta") and os.path.exists(f["ruta"]) and foto_valida(f["ruta"]):
            fotos.append(f)
        elif f.get("ruta") and not os.path.exists(f["ruta"]):
            AVISOS.append(f"No existe la foto {f['ruta']}")
    if len(fotos) > MAX_FOTOS:
        AVISOS.append(f"Hay {len(fotos)} fotos; el anexo incluye las {MAX_FOTOS} primeras.")
        fotos = fotos[:MAX_FOTOS]

    foto_portada = datos.get("foto_portada") or (fotos[0]["ruta"] if fotos else None)
    if foto_portada and not os.path.exists(foto_portada):
        foto_portada = None
    tmpdir = tempfile.mkdtemp()
    hacer_portada = generar_portada_vertical if datos["empresa"] == "vertical" else generar_portada
    ruta_portada = hacer_portada(datos_portada(datos, marca, calc, fechas, foto_portada),
                                 os.path.join(tmpdir, "portada.jpg"))

    # Vertical Projects tiene maqueta propia (scripts/vertical.py) para que no se
    # parezca a Scala. Se registra este módulo como "generar" para que vertical.py
    # use las mismas funciones y la misma lista de avisos.
    clase = Constructor
    if datos["empresa"] == "vertical":
        sys.modules.setdefault("generar", sys.modules[__name__])
        from vertical import ConstructorVertical
        clase = ConstructorVertical

    for intento, ajuste in enumerate(AJUSTES_FOTO):
        doc = clase(datos, marca, calc, fechas, ajuste).montar(ruta_portada, fotos)
        buf = io.BytesIO()
        doc.save(buf)
        tamano = buf.tell()
        if tamano <= LIMITE_BYTES or not fotos or intento == len(AJUSTES_FOTO) - 1:
            break
    if tamano > LIMITE_BYTES:
        AVISOS.append(f"El documento ocupa {tamano / 1e6:.1f} MB: puede superar el límite de subida de la aplicación. "
                      "Quita algunas fotos y vuelve a generarlo.")
    os.makedirs(os.path.dirname(os.path.abspath(salida)), exist_ok=True)
    with open(salida, "wb") as fh:
        fh.write(buf.getvalue())

    print(f"OK  Presupuesto generado: {salida}")
    print(f"    Tamaño: {tamano / 1e6:.2f} MB · fotos en el anexo: {len(fotos)} (lado {ajuste[0]} px, calidad {ajuste[1]})")
    print(f"    PEM: {fmt(calc['pem'])} €", end="")
    if calc["margen_importe"]:
        print(f" · GG y BI {fmt_pct(calc['margen_pct'])} %: {fmt(calc['margen_importe'])} € · Base: {fmt(calc['base'])} €", end="")
    print(f" · IVA {fmt_pct(calc['iva_pct'])} %: {fmt(calc['iva'])} € · TOTAL: {fmt(calc['total'])} €")
    for c in calc["capitulos"]:
        print(f"      {c['codigo']} {c['titulo']}: {fmt(c['subtotal'])} €")
    if calc["pendientes"]:
        print(f"    Partidas pendientes de precio (fuera del total): {calc['pendientes']}")
    for a in AVISOS:
        print(f"AVISO: {a}")


if __name__ == "__main__":
    main()
